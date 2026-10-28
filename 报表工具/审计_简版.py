# 审计_简版.py
# 只做一件事：检查《核算表》里「达成比」这一列，算得对不对。
#
# 用法：python 审计_简版.py "C:\Users\46229\Desktop\某张表.xlsx"
#
# 每行前面带 ①②③ 的标注，是给你看代码用的，不是给电脑看的。

import sys
import openpyxl

# ⓪ Windows 终端默认是 GBK 编码，打印中文和 ✅❌ 会崩，这一行把它改成 UTF-8
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# ① 从命令行拿表格路径，打开文件，选中「1.核算表」这个工作表
path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb["1.核算表"]

# ② 一行一行往下扫。range(4, 32) 的意思是：从第 4 行扫到第 31 行
for row in range(4, 32):
    label = ws.cell(row, 3).value or ws.cell(row, 2).value or ws.cell(row, 1).value
    plan = ws.cell(row, 5).value           # E 列 = 预定金额
    act = ws.cell(row, 7).value            # G 列 = 实际金额
    rate = ws.cell(row, 10).value          # J 列 = 表里写的达成比

    # ③ 跳过没有意义的数据
    if not label:
        continue                           # 这一行没有名称，跳过
    if not isinstance(plan, (int, float)) or plan == 0:
        continue                           # 预定是空的 / 是 "/" / 是 0，没法算，跳过
    if not isinstance(act, (int, float)):
        continue                           # 实际是空的，跳过

    # ④ 自己算一遍：达成比 = 实际 ÷ 预定
    expect = act / plan

    # ⑤ 和表里写的比。rate 可能不是数字（比如 "#DIV/0!"），先排除
    if not isinstance(rate, (int, float)):
        print(f"第 {row:2d} 行 {label}：达成比不是数字（{rate}）")
        continue

    if abs(expect - rate) > 0.005:         # 差距超过 0.5% 才报，容忍四舍五入
        print(f"❌ 第 {row:2d} 行 {label}：应为 {expect:.2%}，表里是 {rate:.2%}")
    else:
        print(f"✅ 第 {row:2d} 行 {label}：{rate:.2%}")

wb.close()
